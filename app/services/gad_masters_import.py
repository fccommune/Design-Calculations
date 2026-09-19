"""
Parses the "GAD Automation Solidworks" master workbook into rows ready to
bulk-load into the 6 lookup tables app/services/gad_globe_lookup.py queries
(ga_globe_table, ga_globe_hookup, ga_crosssec_globe, ga_sheet4_globe,
ga_dim_valve_globe, ga_dim_act_globe) - see app/models.py for the tables
themselves.

The workbook is expected to have (some of) these sheets, matched by name:
    Overall Assy Selection          -> GAGlobeTable
    Hook Up Selection               -> GAGlobeHookUp
    Cross Sec Selection             -> GACrossSecGlobe
    Sheet4 Selection                -> GASheet4Globe
    Dimension Table For Valve       -> GADimValveGlobe
    Dimension Table For Actuator    -> GADimActGlobe
Any other sheets (e.g. a "Tools" index sheet) are ignored. A single-sheet
workbook (like the Sheet4 master's own fully_expanded_configuration.xlsx)
uploaded via parse_single_table() doesn't need its sheet named "Sheet4
Selection" at all - see that function's docstring.
"""
import openpyxl

from app.models import (
    GAGlobeTable,
    GAGlobeHookUp,
    GACrossSecGlobe,
    GASheet4Globe,
    GADimValveGlobe,
    GADimActGlobe,
)


class MasterImportError(Exception):
    """User-facing failure - bad workbook, or a recognized sheet missing an
    expected column."""


# sheet name -> (model, {excel header: model column name})
SHEET_MAP = {
    "Overall Assy Selection": (GAGlobeTable, {
        "Valve_Series": "valve_series",
        "Body_Style": "body_style",
        "Valve_Size ": "valve_size",
        "Rating": "rating",
        "End _Connection": "end_connection",
        "Bonnet _Type": "bonnet_type",
        "Flow_Direction": "flow_direction",
        "Act Series": "actuator_series",
        "Act Type": "actuator_type",
        "Actuator Size": "actuator_size",
        "Traval stop": "traval_stop",
        "H/W": "hw",
        "New Dwg No": "drawing_no",
    }),
    "Hook Up Selection": (GAGlobeHookUp, {
        "Actuator": "actuator",
        "Spring": "spring",
        "Air Fail\naction": "air_fail_action",
        "Positioner": "positioner",
        "Hand \nAuto": "hand_auto",
        "Limit \nSwitch": "limit_switch",
        "Position\nTrans.": "position_trans",
        "Volume\nBooster": "volume_booster",
        "Spool \nValve": "spool_valve",
        "Air\nlock relay": "lock_valve",
        "Solenoid Valve": "solenoid_valve",
        "Schematic\nNo.": "schematic_no",
    }),
    "Cross Sec Selection": (GACrossSecGlobe, {
        "Body Style": "body_style",
        "Size": "size",
        "Bonnet Type": "bonnet_type",
        "Trim Type": "trim_type",
        "Balancing": "balancing",
        "Bal Seal Type": "bal_seal_type",
        "Seat Type": "seat_type",
        "Packing Type": "packing_type",
        "Flow Direction": "flow_direction",
        "Drawing No.": "drawing_no",
    }),
    "Sheet4 Selection": (GASheet4Globe, {
        "Body Type": "body_style",
        "End Connection": "end_connection",
        "Size ": "size",
        "Rating": "rating",
        "Bonnet Type": "bonnet_type",
        "Type": "trim_type",
        "Balancing": "balancing",
        "Flow Direction": "flow_direction",
        "Balance Seal": "bal_seal_type",
        "Seat Type": "seat_type",
        "Packing": "packing_type",
        "DWG No.": "drawing_no",
    }),
    "Dimension Table For Valve": (GADimValveGlobe, {
        "Body_Style": "body_style",
        "End_connection": "end_connection",
        "Bonnet_Type": "bonnet_type",
        "End_finish": "end_finish",
        "series": "series",
        "Size": "size",
        "Rating": "rating",
        "Stem dia": "stem_dia",
        "A": "a",
        "B": "b",
        "C": "c",
        "AR": "ar",
        "Valve Weight": "weight",
    }),
    "Dimension Table For Actuator": (GADimActGlobe, {
        "Actuator_Type": "actuator_type",
        "Actuator Size": "actuator_size",
        "Handwheel": "hand_wheel",
        "Travel Stop": "travel_stop",
        "D": "d",
        "E": "e",
        "F": "f",
        "Weight (Kg)": "weight",
    }),
}

MODELS_BY_TABLE = {model.__tablename__: model for model, _ in SHEET_MAP.values()}

TABLE_LABELS = {
    "ga_globe_table": "Overall Assembly Selection",
    "ga_globe_hookup": "Hook-Up Selection",
    # Legacy/unused - Sheet2 no longer consults this table at all (replaced
    # by ga_sheet4_globe below, now labeled "Cross-Section" itself). Kept
    # only so a direct link to it (e.g. from history) doesn't 404; not
    # listed in the sidebar or the GAD Masters overview page anymore.
    "ga_crosssec_globe": "Cross-Section Selection (legacy, unused)",
    "ga_sheet4_globe": "Cross-Section",
    "ga_dim_valve_globe": "Valve Dimensions",
    "ga_dim_act_globe": "Actuator Dimensions",
}

# table name -> (expected sheet name, {excel header: model column name})
TABLE_SHEET_INFO = {
    model.__tablename__: (sheet_name, header_map)
    for sheet_name, (model, header_map) in SHEET_MAP.items()
}

# table name -> the exact Excel column headers a single-table upload must have.
EXCEL_HEADERS_BY_TABLE = {
    table_name: list(header_map.keys())
    for table_name, (_, header_map) in TABLE_SHEET_INFO.items()
}


# Columns that legitimately hold a numeric weight (kept as None/NULL when
# blank so the lookup's `weight or 0` fallback applies) - every other mapped
# column is a NOT NULL string column, so a blank cell there becomes "" to
# match how app.services.gad_globe_lookup.cell() reads a blank BOM cell as
# "" rather than None.
_NUMERIC_FIELDS = {"weight"}


def _cell(value):
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return str(value).strip()


def _parse_sheet(worksheet, header_map):
    rows_iter = worksheet.iter_rows(values_only=True)
    try:
        raw_headers = next(rows_iter)
    except StopIteration:
        return []

    col_index = {}
    for i, raw in enumerate(raw_headers):
        name = "" if raw is None else str(raw)
        if name in header_map:
            col_index[header_map[name]] = i

    missing = set(header_map.values()) - set(col_index.keys())
    if missing:
        raise MasterImportError(
            f"Sheet '{worksheet.title}' is missing expected column(s): {', '.join(sorted(missing))}"
        )

    rows = []
    for raw_row in rows_iter:
        if raw_row is None or all(v is None for v in raw_row):
            continue
        record = {}
        for field, idx in col_index.items():
            value = raw_row[idx] if idx < len(raw_row) else None
            cell_value = _cell(value)
            if cell_value is None and field not in _NUMERIC_FIELDS:
                cell_value = ""
            record[field] = cell_value
        if not any(record.values()):
            continue
        rows.append(record)
    return rows


def parse_master_workbook(file_stream) -> dict:
    """Returns {table_name: [row dict, ...]} for every recognized sheet found
    in the workbook (missing sheets are simply absent from the result).
    Raises MasterImportError if the workbook can't be read, if none of the
    expected sheets are present, or if a present sheet is missing an
    expected column."""
    try:
        workbook = openpyxl.load_workbook(file_stream, data_only=True, read_only=True)
    except Exception as exc:
        raise MasterImportError(f"Failed to load Excel: {exc}") from exc

    result = {}
    for sheet_name, (model, header_map) in SHEET_MAP.items():
        if sheet_name not in workbook.sheetnames:
            continue
        rows = _parse_sheet(workbook[sheet_name], header_map)
        result[model.__tablename__] = rows

    workbook.close()

    if not result:
        raise MasterImportError(
            "No recognized master sheets found. Expected one or more of: "
            + ", ".join(SHEET_MAP.keys())
        )
    return result


def parse_single_table(file_stream, table_name: str) -> list:
    """Parses a workbook meant for just one master table - e.g. the table's
    own sheet exported to its own file. Reads the sheet named after that
    table if present, otherwise falls back to the workbook's first sheet
    (so a plain single-sheet export works too), and requires that sheet to
    have every column the table expects."""
    if table_name not in TABLE_SHEET_INFO:
        raise MasterImportError(f"Unknown master table '{table_name}'.")
    sheet_name, header_map = TABLE_SHEET_INFO[table_name]

    try:
        workbook = openpyxl.load_workbook(file_stream, data_only=True, read_only=True)
    except Exception as exc:
        raise MasterImportError(f"Failed to load Excel: {exc}") from exc

    worksheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else workbook.worksheets[0]
    rows = _parse_sheet(worksheet, header_map)
    workbook.close()
    return rows
